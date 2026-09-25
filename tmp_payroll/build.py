import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open("rows.json"))
cy = d["cycle"]
wb = Workbook()

FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE  = Font(name=FONT, bold=True, size=14)
SUB    = Font(name=FONT, size=10, color="555555")
BODY   = Font(name=FONT, size=10)
BOLD   = Font(name=FONT, size=10, bold=True)
ASKF   = PatternFill("solid", fgColor="FFF2CC")   # the columns Cyrine fills
MONEY  = PatternFill("solid", fgColor="FCE4D6")   # pay is affected
THIN   = Border(*[Side(style="thin", color="D9D9D9")]*4)
TOP    = Alignment(vertical="top", wrap_text=True)

def head(ws, cols, row=1):
    for i, (label, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = H_FILL, H_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30

# ---------------------------------------------------------------- Sheet 1
ws = wb.active
ws.title = "Answer these"
ws["A1"] = f"{cy['label']} — {len(d['questions'])} questions for you"
ws["A1"].font = TITLE
ws["A2"] = (f"Period {cy['start']} to {cy['end']}. Only you can answer these — they are days where the record says "
            "somebody did not come in, or left long before the end, and nothing in the data can tell us whether "
            "that is right. Everything else is either being corrected at this end ('We are correcting these') or "
            "needs a yes/no click ('Overtime waiting'). "
            "Highest impact on pay first. Fill in the two yellow columns and send this file back; "
            "nothing here needs to be fixed in the system by you unless a row says so.")
ws["A2"].font = SUB
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:L2"); ws.row_dimensions[2].height = 42

cols = [("#",4),("Staff",22),("Date",11),("What it is",18),
        ("What we need you to check",46),("What the system has now",26),
        ("Shift published",24),("Clocked in–out",22),
        ("What happens to pay",34),("➜ YOUR ANSWER",26),("➜ Your note",28),("Then do this in the system",38)]
head(ws, cols, row=4)
ws.freeze_panes = "A5"

r = 5
for q in d["questions"]:
    vals = [q["rank"], q["staff"], q["date"], q["kind"], q["check"],
            q["system"], q["roster"], q["punch"], q["impact"], "", "", q["fix"]]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = BOLD if i == 2 else BODY
        c.alignment = TOP
        c.border = THIN
        if i in (10, 11):
            c.fill = ASKF
        if i == 12:
            c.font = Font(name=FONT, size=9, color="666666")
        if i == 9 and "Deducts" in str(q["impact"]):
            c.fill = MONEY
    # a dropdown of the sensible answers, so the reply comes back in words we can act on
    opts = q["answer_options"].replace('"', "'")
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(o.strip() for o in opts.split("/")),
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv); dv.add(ws.cell(row=r, column=10))
    ws.row_dimensions[r].height = 46
    r += 1

# ---------------------------------------------------------------- Sheet 2
ws2 = wb.create_sheet("Before you compute")
ws2["A1"] = "Do this on the night of 25 September, before the payroll is computed"
ws2["A1"].font = TITLE
ws2["A3"] = ("1.  Wait until the last shift of 25 September has finished and everyone has clocked out.\n"
             "2.  Run 'Sync from OS Attendance' for Dubai so the DTR picks up those clock-outs.\n"
             "3.  Then recompute the payroll adjustments.\n\n"
             "Why: the figures currently stored were computed on 24 September, before your corrections, "
             "so they have to be recomputed anyway. And the "
             f"{len(d['open_punchout_0925'])} people below still had no clock-out for 25 September when this "
             "file was made. If the payroll is computed after 25 September without a re-sync, each of them is "
             "charged a one-hour admin fee for a missing punch — for a day that was simply still in progress.")
ws2["A3"].font = BODY
ws2["A3"].alignment = Alignment(wrap_text=True, vertical="top")
ws2.merge_cells("A3:D3"); ws2.row_dimensions[3].height = 150
ws2.column_dimensions["A"].width = 34
for col in "BCD": ws2.column_dimensions[col].width = 20

ws2["A5"] = f"No clock-out yet for 25 Sep ({len(d['open_punchout_0925'])} people)"
ws2["A5"].font = Font(name=FONT, bold=True, size=11)
head(ws2, [("Staff",34),("Clocked out now?",20),("",20),("",20)], row=6)
for i, n in enumerate(d["open_punchout_0925"]):
    c = ws2.cell(row=7+i, column=1, value=n); c.font = BODY; c.border = THIN
    c2 = ws2.cell(row=7+i, column=2, value=""); c2.fill = ASKF; c2.border = THIN
ws2.freeze_panes = "A7"

# ---------------------------------------------------------------- Sheet 3
ws3 = wb.create_sheet("Already clear")
ws3["A1"] = f"Attendance is clear — {len(d['clear'])} of the 61 people"
ws3["A1"].font = TITLE
ws3["A2"] = ("Their attendance was checked against the live DTR and needs no answer from you. Some do carry a "
             "deduction (lateness, undertime, a missed punch), but each one is backed by the clock record, "
             "so there is nothing to confirm. This covers ATTENDANCE ONLY — ten of the people on this list "
             "also have an overtime claim nobody has approved or rejected yet. Those are on the "
             "'Overtime waiting' tab and still need doing.")
ws3["A2"].font = SUB
ws3["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws3.merge_cells("A2:C2"); ws3.row_dimensions[2].height = 46
head(ws3, [("Staff",30),("Staff",30),("Staff",30)], row=4)
names = d["clear"]
per = (len(names) + 2) // 3
for col in range(3):
    chunk = names[col*per:(col+1)*per]
    for i, n in enumerate(chunk):
        c = ws3.cell(row=5+i, column=col+1, value=n); c.font = BODY; c.border = THIN
ws3.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 1b
ws7 = wb.create_sheet("We are correcting these")
ws7["A1"] = f"Being corrected at this end — {len(d['we_fix'])} items, no answer needed"
ws7["A1"].font = TITLE
ws7["A2"] = ("These are decided by the record itself, so they are not questions. They are listed so you can see "
             "what is changing and say so if any of it is wrong. If you do not reply to this tab, these go ahead.")
ws7["A2"].font = SUB
ws7["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws7.merge_cells("A2:E2"); ws7.row_dimensions[2].height = 40
head(ws7, [("Staff",24),("Date",13),("What the record shows",34),("What we are doing",56),("➜ Object?",16)], row=4)
for i, f in enumerate(d["we_fix"]):
    for j, v in enumerate((f["staff"], f["date"], f["finding"], f["action"], "")):
        c = ws7.cell(row=5+i, column=j+1, value=v)
        c.font = BOLD if j == 0 else BODY; c.alignment = TOP; c.border = THIN
        if j == 4: c.fill = ASKF
    ws7.row_dimensions[5+i].height = 46
ws7.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 1c
ws8 = wb.create_sheet("Charged, not queried")
ws8["A1"] = f"Charged and left as it stands — {len(d['charged_not_queried'])} days"
ws8["A1"].font = TITLE
ws8["A2"] = ("Small lateness and early finishes where the clock is unambiguous AND the same person clocks in on "
             "time at that same start time on their other days. We are not asking about these — they are listed "
             "so you can scan them in a minute and tell us if any one of them is wrong. The largest is under a "
             "fiftieth of a month's pay. If you say nothing, they stand.")
ws8["A2"].font = SUB
ws8["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws8.merge_cells("A2:E2"); ws8.row_dimensions[2].height = 54
head(ws8, [("Date",12),("Staff",24),("What it is",20),("Shift published",26),("Clocked in–out",24)], row=4)
for i, z in enumerate(d["charged_not_queried"]):
    for j, v in enumerate((z["date"], z["staff"], z["kind"], z["roster"], z["punch"])):
        c = ws8.cell(row=5+i, column=j+1, value=v)
        c.font = BOLD if j == 1 else BODY; c.border = THIN; c.alignment = TOP
ws8.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 3b
ws5 = wb.create_sheet("Overtime waiting")
ws5["A1"] = f"Overtime nobody has approved or rejected — {len(d['overtime_pending'])} claims"
ws5["A1"].font = TITLE
ws5["A2"] = ("This is money owed TO staff, not a deduction, which is why it is easy to miss — nothing on the "
             "payroll screen complains about it. These claims are inside the September period and are still "
             "sitting as 'pending', so they pay nothing. Approve or reject each one in Overtime Requests "
             "before the payroll is computed. An approved claim also has to be added to payroll — approving "
             "alone does not pay it.")
ws5["A2"].font = SUB
ws5["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws5.merge_cells("A2:E2"); ws5.row_dimensions[2].height = 60
head(ws5, [("Date",12),("Staff",26),("Claimed",12),("➜ Approve / Reject",22),("➜ Your note",34)], row=4)
for i, o in enumerate(d["overtime_pending"]):
    for j, v in enumerate((o["date"], o["staff"], f"{o['min']} min", "", "")):
        c = ws5.cell(row=5+i, column=j+1, value=v)
        c.font = BOLD if j == 1 else BODY; c.alignment = TOP; c.border = THIN
        if j >= 3: c.fill = ASKF
    dv = DataValidation(type="list", formula1='"Approve,Reject"', allow_blank=True, showDropDown=False)
    ws5.add_data_validation(dv); dv.add(ws5.cell(row=5+i, column=4))
r5 = 5 + len(d["overtime_pending"]) + 2
ws5.cell(row=r5, column=1, value="Older claims, outside this period").font = Font(name=FONT, bold=True, size=11)
ws5.cell(row=r5+1, column=1, value=(
    "These three are still pending but fall in a period that is already closed, so they cannot be paid by "
    "recomputing September. Please say what should happen to them — they will not resolve themselves.")
).font = BODY
ws5.cell(row=r5+1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws5.merge_cells(start_row=r5+1, start_column=1, end_row=r5+1, end_column=5)
ws5.row_dimensions[r5+1].height = 34
for i, o in enumerate(d["overtime_stranded"]):
    for j, v in enumerate((o["date"], o["staff"], f"{o['min']} min", "", "")):
        c = ws5.cell(row=r5+2+i, column=j+1, value=v)
        c.font = BOLD if j == 1 else BODY; c.border = THIN
        if j >= 3: c.fill = ASKF
ws5.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 3c
ws6 = wb.create_sheet("Clock-outs to repair")
ws6["A1"] = f"Days recorded as zero hours worked — {len(d['zero_hours'])} rows"
ws6["A1"].font = TITLE
ws6["A2"] = ("On these days the person clocked in, and the clock-out landed on the following day, so the "
             "system credited zero hours. NONE of these change September pay — they are all monthly staff or "
             "days outside the hourly window, and nothing is deducted. They are listed so the DTR you sign "
             "off is not carrying days that say nobody worked. Repair them when there is time; they do not "
             "hold up the payroll. Four of them are Bibek BK on four nights in a row, which suggests the "
             "night shift is the cause rather than the person.")
ws6["A2"].font = SUB
ws6["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws6.merge_cells("A2:D2"); ws6.row_dimensions[2].height = 74
head(ws6, [("Date",12),("Staff",26),("Clocked in → out",30),("➜ Fixed?",14)], row=4)
for i, z in enumerate(d["zero_hours"]):
    for j, v in enumerate((z["date"], z["staff"], z["punch"], "")):
        c = ws6.cell(row=5+i, column=j+1, value=v)
        c.font = BOLD if j == 1 else BODY; c.border = THIN
        if j == 3: c.fill = ASKF
ws6.freeze_panes = "A5"

# ---------------------------------------------------------------- Sheet 4
ws4 = wb.create_sheet("What we checked")
ws4["A1"] = "What was checked, and what was found to be already correct"
ws4["A1"].font = TITLE
notes = [
 ("Split shifts you re-typed", "All 21 rows now carry both segments (e.g. 12:00-16:00 & 18:00-22:00) and record no undertime at all. Regular hours 8 on twenty of them and 7.99 on one, which is rounding, not a short day. Correct."),
 ("Annual leave", "147 leave days across 9 people are flagged as annual leave in the DTR, so none of them is deducted. Correct."),
 ("Nishan Nepal 15 Sep", "Now shows as a rest day. Correct."),
 ("Man Bahadur B K 21 Sep", "Shift is published and the DTR has it. Correct. His clock-in time on that day is one of the questions on the 'Answer these' tab."),
 ("'Published as DAY OFF but DTR says working day'", "106 rows on the previous file. NONE of these affect pay — Dubai only deducts a day that is marked absent without pay, and the day type is not part of the calculation. Please ignore them. They were listed first on the previous file, which was our mistake."),
 ("Everything that moves money", "Every rule the Dubai payroll applies was checked for all 61 people: absence, lateness, undertime, break excess, missing punches, the 5% monthly late penalty, night premium and approved overtime. The questions on the first tab are every case where the answer is not already in the clock record."),
 ("Christian Baria and Shushma Kumari", "Both have left the company and their final pay is settled, so neither appears on any list here. Their salary records are still switched on, so whoever builds the September run must leave them out of it. The run has not been created yet, so nothing has been paid twice."),
 ("Yogesh Bashyal", "He had used up his annual leave and did not come back, so he is treated as having resigned. The annual leave that falls inside this period is 26 to 31 August and is paid; 1 to 25 September is not. The DTR currently flags 1 to 10 September as PAID annual leave, which is wrong and needs clearing — but do NOT mark those days absent without pay, because the part-month adjustment already covers the whole of 1 to 25 September and marking them as well would deduct twice. His salary record is still switched on, so he must also be left out of the September run."),
 ("Worked during annual leave", "Nobody. Nothing to correct."),
 ("Worked on a rest day", "Nobody. Nothing to correct."),
 ("Anyone missing from the calculation", "Nobody. All 61 people with attendance have an active salary record."),
 ("Part-time (hourly) staff", "8 people are paid by the hour: Bijien Mijar, Dipesh Thapa, Kelvin Gurung, Mahima Pansilu Dadallage, Padam Bahadur K C, Pukar K C, Raman Miya, Rubash Khadka. Lateness, absence and undertime carry no penalty for them, so their records are not on the question list even where a day looks odd."),
 ("Jeffril Marcos Vergara, 24 Sep", "He worked this day with no shift published. A shift of 15:00-24:00 was published while this review was being written, so it is no longer an open question. He clocked 12:56-22:08, which is about two hours earlier than the published shift at both ends, but he is recorded as not late and with a full eight hours, so nothing is deducted and there is nothing for you to do."),
 ("Staff who joined or left mid-period", "Man Bahadur B K joined 21 Sep, Sangita Giri K C joined 19 Sep, Dipesh Thapa joined 12 Sep, Chandra Gurung last worked 11 Sep, Sota Horii resigned 12 Sep, Bibek Tamang starts 27 Sep. All six already have a part-month adjustment recorded."),
]
head(ws4, [("Item",44),("Result",96)], row=3)
for i, (a, b) in enumerate(notes):
    for j, v in enumerate((a, b)):
        c = ws4.cell(row=4+i, column=j+1, value=v)
        c.font = BOLD if j == 0 else BODY
        c.alignment = TOP; c.border = THIN
    ws4.row_dimensions[4+i].height = 44
ws4.freeze_panes = "A4"

out = "Dubai_September_payroll_open_questions_20260925.xlsx"
wb.save(out)
print("wrote", out)
