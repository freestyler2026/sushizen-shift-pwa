import json, re, sys
from openpyxl import load_workbook
T = json.load(open("truth.json"))
wb = load_workbook("Dubai_September_payroll_open_questions_20260925.xlsx", data_only=True)
fail, checked = [], 0
def eq(label, got, want):
    global checked; checked += 1
    if str(got).strip() != str(want).strip():
        fail.append(f"{label}\n      file : {got!r}\n      truth: {want!r}")

# --- Sheet 1 : every evidence cell against production -------------------
ws = wb["Answer these"]
hdr = [c.value for c in ws[4]]
assert hdr[:4] == ["#","Staff","Date","What it is"], hdr[:4]
i_ros, i_pun = hdr.index("Shift published"), hdr.index("Clocked in–out")
i_sys, i_ans = hdr.index("What the system has now"), hdr.index("➜ YOUR ANSWER")
seen, ranks = set(), []
for r in range(5, ws.max_row + 1):
    if ws.cell(r, 2).value is None: continue
    staff, date = ws.cell(r, 2).value, ws.cell(r, 3).value
    k = f"{staff}|{date}"
    ranks.append(ws.cell(r, 1).value)
    if k not in T["ev"]:
        fail.append(f"row {r}: {k} is not in the verified production set"); continue
    seen.add(k); t = T["ev"][k]
    eq(f"row {r} {k} roster", ws.cell(r, i_ros+1).value, t["ros"])
    eq(f"row {r} {k} punch",  ws.cell(r, i_pun+1).value, t["punch"])
    # the "system says" text must not contradict the numbers
    sysx = str(ws.cell(r, i_sys+1).value)
    checked += 1
    if t["late"] and f'{t["late"]} min late' not in sysx:
        fail.append(f'row {r} {k}: says {sysx!r} but production late={t["late"]}')
    checked += 1
    if t["awp"] and "absent without pay" not in sysx:
        fail.append(f'row {r} {k}: production is absent-without-pay, cell says {sysx!r}')
    checked += 1
    if t.get("brk") and f'{t["brk"]} min' not in sysx:
        fail.append(f'row {r} {k}: production break={t["brk"]} not stated in {sysx!r}')
    checked += 1
    if t["und"] and f'{t["und"]} min undertime' not in sysx:
        fail.append(f'row {r} {k}: production undertime={t["und"]} not stated in {sysx!r}')
    checked += 1
    if not str(ws.cell(r, hdr.index("Then do this in the system")+1).value or "").strip():
        fail.append(f"row {r} {k}: the 'then do this' cell is empty")
    checked += 1
    if ws.cell(r, i_ans+1).value not in (None, ""):
        fail.append(f"row {r}: the answer column is not blank")
missing = set(T["ev"]) - seen
if missing: fail.append(f"questions missing from the file: {sorted(missing)}")
checked += 1
if ranks != list(range(1, len(ranks)+1)): fail.append(f"# column is not 1..n: {ranks}")

# --- Sheet 2 : the open punch-outs --------------------------------------
ws2 = wb["Before you compute"]
names2 = [ws2.cell(r,1).value for r in range(7, ws2.max_row+1) if ws2.cell(r,1).value]
eq("open punch-out list", sorted(names2), sorted(T["open25"]))
eq("open punch-out count in the text", str(len(T["open25"])) in str(ws2["A3"].value), "True")

# --- Sheet 3 : cleared ---------------------------------------------------
ws3 = wb["Already clear"]
clear = [ws3.cell(r,c).value for r in range(5, ws3.max_row+1) for c in (1,2,3) if ws3.cell(r,c).value]
askers = {k.split("|")[0] for k in T["ev"]} - {"Four people"}
askers |= {"Jheymar Fabros","Manisha Tamang Gurung","Sita Gurmachhan","Aris Jhon De Ocampo"}
checked += 1
overlap = askers & set(clear)
if overlap: fail.append(f"these people are on BOTH lists: {sorted(overlap)}")
eq("no duplicate names in the cleared list", len(clear), len(set(clear)))
# --- Sheet 4 : the claims ------------------------------------------------
ws4 = wb["What we checked"]
txt = " ".join(str(ws4.cell(r,c).value) for r in range(4, ws4.max_row+1) for c in (1,2))
eq("annual-leave days claim", f'{T["annual_days"]} leave days' in txt, "True")
eq("annual-leave people claim", f'{T["annual_people"]} people' in txt, "True")
eq("hourly staff named", all(n in txt for n in T["hourly"]), "True")
eq("hourly staff count", "8 people are paid by the hour" in txt, "True")
eq("total staff claim", f'All {T["dtr_staff"]} people' in txt, "True")
# 21 asked + 38 clear + 2 settled outside this payroll = the 61 people in the DTR
eq("asked + clear + settled = everyone in the DTR",
   len(clear) + len(askers) + len(T["settled_elsewhere"]), T["dtr_staff"])
checked += 1
for n in T["settled_elsewhere"]:
    if n in clear or n in askers:
        fail.append(f"{n} left the company and is settled elsewhere, but still appears on a list")
ws4txt = txt
checked += 1
if not all(n in ws4txt for n in T["settled_elsewhere"]):
    fail.append("a leaver is not named anywhere in the file")
# the substance, not just the name: each note must carry the instruction that matters
MUST = {
  "Yogesh Bashyal": ["26 to 31 August", "1 to 25 September",
                     "do NOT mark those days absent without pay", "left out of the September run"],
  "Christian Baria and Shushma Kumari": ["leave them out of it", "final pay is settled"],
}
for row in range(4, ws4.max_row + 1):
    label = str(ws4.cell(row, 1).value or "")
    if label in MUST:
        body = str(ws4.cell(row, 2).value or "")
        for phrase in MUST[label]:
            checked += 1
            if phrase not in body:
                fail.append(f'the note for "{label}" no longer says: {phrase!r}')
# a cross-reference by row number goes stale the moment a question is removed
import re as _re
for _r in range(1, ws4.max_row+1):
    for _c in (1,2):
        _v = str(ws4.cell(_r,_c).value or "")
        checked += 1
        if _re.search(r"\brow\s+\d+\b", _v, _re.I):
            fail.append(f"What we checked!{ws4.cell(_r,_c).coordinate} points at a row number, which goes stale: {_v[:70]!r}")
for label in MUST:
    checked += 1
    if not any(str(ws4.cell(r,1).value or "") == label for r in range(4, ws4.max_row+1)):
        fail.append(f'the note "{label}" is missing entirely')



# --- Overtime sheet against production --------------------------------
ws5 = wb["Overtime waiting"]
OT = T["overtime_pending"]; ST = T["overtime_stranded"]
got = []
for r in range(5, 5+len(OT)):
    got.append((ws5.cell(r,1).value, ws5.cell(r,2).value, ws5.cell(r,3).value))
    checked += 1
    if ws5.cell(r,4).value not in (None,""): fail.append(f"overtime row {r}: decision pre-filled")
eq("pending overtime rows", got, [(o["date"], o["staff"], f'{o["min"]} min') for o in OT])
tail = [(ws5.cell(r,1).value, ws5.cell(r,2).value, ws5.cell(r,3).value)
        for r in range(5+len(OT)+4, 5+len(OT)+4+len(ST))]
eq("stranded overtime rows", tail, [(o["date"], o["staff"], f'{o["min"]} min') for o in ST])
# the cleared list must no longer claim those people need nothing
ot_names = {o["staff"] for o in OT}
both = ot_names & set(clear)
checked += 1
if both and "Overtime waiting" not in str(ws3["A2"].value):
    fail.append(f"cleared sheet still claims nothing is needed for {sorted(both)}")

# --- Clock-outs to repair ------------------------------------------------
ws6 = wb["Clock-outs to repair"]
zrows = [(ws6.cell(r,1).value, ws6.cell(r,2).value, ws6.cell(r,3).value)
         for r in range(5, 5+len(T["zero_hours"]))]
eq("zero-hour rows", zrows, [(z["date"], z["staff"], z["punch"]) for z in T["zero_hours"]])
checked += 1
if "NONE of these change September pay" not in str(ws6["A2"].value):
    fail.append("zero-hours sheet does not say these do not affect pay")

# --- no salary figures anywhere ------------------------------------------
money = []
for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            v = str(c.value or "")
            if re.search(r"\bAED\b|\b\d{3,5}\.\d{2}\b", v): money.append(f"{sh.title}!{c.coordinate}: {v[:70]}")
checked += 1
if money: fail.append("money figures found in the file:\n      " + "\n      ".join(money))

print(f"{checked} checks run")
if fail:
    print(f"\n{len(fail)} FAILED:\n")
    for f in fail: print("  ✗", f)
    sys.exit(1)
print("\nall checks passed")
